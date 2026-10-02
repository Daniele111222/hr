from io import BytesIO

from openpyxl import load_workbook


def test_tax_auxiliary_preserves_headers_text_and_unknown_fields_without_tax_calculation():
    from paylite.excel.tax_export import TEMPLATE_PATH, render_tax

    template = load_workbook(TEMPLATE_PATH)
    row = {
        "subject_id": 1,
        "subject_name": "测试主体",
        "employee_id": 1,
        "payroll_batch_id": 2,
        "batch_type": "normal",
        "batch_no": 2,
        "correction_of_batch_id": 1,
        "snapshot": {"name": "=员工", "employee_no": "001", "id_number": "001101199001011234"},
        "amounts": {"gross": "1000.00", "untaxed_amount": "900.00"},
        "items": [
            {"code": "social_employee_pension", "amount": "60.00"},
            {"code": "housing_employee", "amount": "40.00"},
        ],
    }
    supplement = {
        **row,
        "payroll_batch_id": 3,
        "batch_type": "supplement",
        "batch_name": "=补发原因",
        "correction_of_batch_id": None,
        "amounts": {"gross": "50.00", "untaxed_amount": "50.00"},
        "items": [{"code": "supplement", "amount": "50.00"}],
    }
    book = load_workbook(BytesIO(render_tax({"period": "2026-09", "records": [supplement, row]})))
    sheet = book["Sheet1"]
    assert book.sheetnames == ["Sheet1"]
    assert [c.value for c in sheet[1]] == [c.value for c in template.active[1]]
    assert sheet["G1"].value == "本期收入"
    assert sheet["A2"].value == "001" and sheet["A2"].data_type == "s"
    assert sheet["B2"].data_type == "s"
    assert sheet["D2"].value == "001101199001011234"
    assert sheet["D2"].number_format == "@"
    assert [sheet[f"{c}2"].value for c in ("G", "I", "L")] == [1000, 60, 40]
    assert sheet["G3"].value == 50
    for n in (2, 3):
        for column in ("C", "E", "F", "H", "J", "K", "M", "AB", "AC", "AE", "AF"):
            assert sheet[f"{column}{n}"].value is None
    assert sheet["I3"].value is None and sheet["L3"].value is None
    assert "未扣个税金额 900.00" in sheet["AG2"].value
    assert "整批更正" in sheet["AG2"].value
    assert "独立补发" in sheet["AG3"].value and "=补发原因" in sheet["AG3"].value
    assert "未完成税务申报" in sheet["AG2"].value
    assert sheet["A1"]._style == template.active["A1"]._style
    assert not book._external_links and not book.defined_names
    assert not any(c.data_type in {"e", "f"} for r in sheet for c in r)
