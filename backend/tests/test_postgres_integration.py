import os
from collections.abc import Generator
from concurrent.futures import ThreadPoolExecutor
from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path

import pytest
from alembic import command
from alembic.config import Config
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.engine import URL, make_url
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

import paylite.services.payroll_trial as payroll_trial_service
from paylite.api.deps import get_db
from paylite.application import create_app
from paylite.config import get_settings
from paylite.db.models import (
    AttendanceRecord,
    AttendanceRule,
    City,
    Company,
    CorrectionBatch,
    Department,
    Employee,
    EmployeeAssignment,
    EmployeeBase,
    EmployeeSalary,
    HousingFundRule,
    ImportBatch,
    ImportRow,
    PayrollBatch,
    PayrollPeriod,
    PayrollRecord,
    PayrollTrialRun,
    PerformanceRecord,
    SocialSecurityItemRule,
    SocialSecurityRule,
    Subject,
    SubjectDepartment,
)
from paylite.excel.attendance_template import make_template as make_attendance_template
from paylite.excel.employee_template import make_template
from paylite.excel.performance_template import make_template as make_performance_template
from paylite.services.payroll_workbench import PayrollWorkbenchError, create_batch

TEST_DATABASE_URL = os.getenv("PAYLITE_TEST_DATABASE_URL")
ALLOW_TEST_DATABASE_RESET = os.getenv("PAYLITE_ALLOW_TEST_DATABASE_RESET") == "1"


def validate_test_database_url(database_url: str) -> URL:
    parsed = make_url(database_url)
    database_name = parsed.database or ""
    if database_name.lower() in {"", "postgres", "paylite"} or "test" not in database_name.lower():
        raise RuntimeError(
            "PAYLITE_TEST_DATABASE_URL must target a dedicated database whose name contains 'test'"
        )
    if not ALLOW_TEST_DATABASE_RESET:
        raise RuntimeError(
            "Set PAYLITE_ALLOW_TEST_DATABASE_RESET=1 to allow integration-test database reset"
        )
    return parsed


pytestmark = pytest.mark.skipif(
    not TEST_DATABASE_URL,
    reason="PAYLITE_TEST_DATABASE_URL is required for PostgreSQL integration tests",
)


@pytest.fixture(scope="session")
def postgres_engine():
    assert TEST_DATABASE_URL is not None
    validate_test_database_url(TEST_DATABASE_URL)
    os.environ["PAYLITE_DATABASE_URL"] = TEST_DATABASE_URL
    get_settings.cache_clear()

    backend_root = Path(__file__).parents[1]
    alembic_config = Config(str(backend_root / "alembic.ini"))
    command.upgrade(alembic_config, "head")

    engine = create_engine(TEST_DATABASE_URL, pool_pre_ping=True)
    yield engine

    if engine.url.database != "paylite":
        command.downgrade(alembic_config, "base")
    engine.dispose()


@pytest.fixture
def db_session(postgres_engine) -> Generator[Session]:
    with Session(postgres_engine) as session:
        yield session
        session.rollback()


@pytest.fixture
def api_client(postgres_engine):
    application = create_app()

    def override_get_db() -> Generator[Session]:
        with Session(postgres_engine) as session:
            try:
                yield session
            except Exception:
                session.rollback()
                raise

    application.dependency_overrides[get_db] = override_get_db
    with TestClient(application) as client:
        yield client
    application.dependency_overrides.clear()


def create_base_records(
    session: Session, suffix: str = ""
) -> tuple[Company, City, Subject, PayrollPeriod, Employee]:
    company = Company(code=f"ACME{suffix}", name="Acme")
    city = City(code=f"BJ{suffix}", name="北京")
    session.add_all([company, city])
    session.flush()

    subject = Subject(company_id=company.id, code=f"MAIN{suffix}", name="主主体")
    period = PayrollPeriod(
        year=2026,
        month=6,
        period_start=date(2026, 6, 1),
        period_end=date(2026, 6, 30),
    )
    employee = Employee(
        company_id=company.id,
        id_number=f"11010119900101123X{suffix}",
        employee_no=f"E001{suffix}",
        name="测试员工",
        employee_type="employee",
        formal_status=True,
        probation_status="confirmed",
        hire_date=date(2020, 1, 1),
    )
    session.add_all([subject, period, employee])
    session.flush()
    return company, city, subject, period, employee


def test_id_number_is_unique_text(db_session: Session) -> None:
    company, _, _, _, _ = create_base_records(db_session)
    duplicate = Employee(
        company_id=company.id,
        id_number="11010119900101123X",
        employee_no="E002",
        name="重复员工",
        employee_type="employee",
        hire_date=date(2026, 1, 1),
    )
    db_session.add(duplicate)

    with pytest.raises(IntegrityError):
        db_session.flush()


def test_normal_batch_is_unique_but_supplement_is_independent(db_session: Session) -> None:
    _, _, subject, period, _ = create_base_records(db_session)
    normal = PayrollBatch(
        subject_id=subject.id,
        payroll_period_id=period.id,
        batch_type="normal",
        status="draft",
    )
    supplement = PayrollBatch(
        subject_id=subject.id,
        payroll_period_id=period.id,
        batch_type="supplement",
        status="draft",
    )
    db_session.add_all([normal, supplement])
    db_session.flush()

    duplicate_normal = PayrollBatch(
        subject_id=subject.id,
        payroll_period_id=period.id,
        batch_type="normal",
        batch_no=2,
        status="draft",
    )
    db_session.add(duplicate_normal)
    with pytest.raises(IntegrityError):
        db_session.flush()


def test_duplicate_import_row_location_is_rejected(db_session: Session) -> None:
    _, _, _, period, employee = create_base_records(db_session)
    import_batch = ImportBatch(
        payroll_period_id=period.id,
        import_type="attendance",
        original_filename="attendance.xlsx",
        file_sha256="a" * 64,
        template_version="v1",
    )
    db_session.add(import_batch)
    db_session.flush()

    db_session.add_all(
        [
            ImportRow(
                import_batch_id=import_batch.id,
                sheet_name="Sheet1",
                source_row_number=2,
                validation_status="valid",
                raw_data={"employee": employee.employee_no},
            ),
            ImportRow(
                import_batch_id=import_batch.id,
                sheet_name="Sheet1",
                source_row_number=2,
                validation_status="valid",
                raw_data={"employee": employee.employee_no},
            ),
        ]
    )
    with pytest.raises(IntegrityError):
        db_session.flush()


def test_duplicate_attendance_for_period_is_rejected(db_session: Session) -> None:
    _, _, _, period, employee = create_base_records(db_session)
    first_batch = ImportBatch(
        payroll_period_id=period.id,
        import_type="attendance",
        original_filename="attendance-1.xlsx",
        file_sha256="b" * 64,
        template_version="v1",
    )
    second_batch = ImportBatch(
        payroll_period_id=period.id,
        import_type="attendance",
        original_filename="attendance-2.xlsx",
        file_sha256="c" * 64,
        template_version="v1",
    )
    db_session.add_all([first_batch, second_batch])
    db_session.flush()

    db_session.add(
        AttendanceRecord(
            import_batch_id=first_batch.id,
            payroll_period_id=period.id,
            employee_id=employee.id,
            expected_work_days=Decimal("20"),
        )
    )
    db_session.flush()
    db_session.add(
        AttendanceRecord(
            import_batch_id=second_batch.id,
            payroll_period_id=period.id,
            employee_id=employee.id,
            expected_work_days=Decimal("20"),
        )
    )

    with pytest.raises(IntegrityError):
        db_session.flush()


def test_rule_is_unique_by_city_and_effective_start(db_session: Session) -> None:
    _, city, _, _, _ = create_base_records(db_session)
    db_session.add(
        SocialSecurityRule(
            city_id=city.id,
            effective_from=date(2026, 1, 1),
            version="2026-01",
        )
    )
    db_session.flush()
    db_session.add(
        SocialSecurityRule(
            city_id=city.id,
            effective_from=date(2026, 1, 1),
            version="2026-01-revision",
        )
    )

    with pytest.raises(IntegrityError):
        db_session.flush()


def test_payroll_record_keeps_snapshot_and_correction_link(db_session: Session) -> None:
    _, _, subject, period, employee = create_base_records(db_session)
    original = PayrollBatch(
        subject_id=subject.id,
        payroll_period_id=period.id,
        batch_type="normal",
        status="locked",
    )
    replacement = PayrollBatch(
        subject_id=subject.id,
        payroll_period_id=period.id,
        batch_type="supplement",
        batch_no=1,
        status="draft",
    )
    db_session.add_all([original, replacement])
    db_session.flush()

    record = PayrollRecord(
        payroll_batch_id=original.id,
        payroll_period_id=period.id,
        subject_id=subject.id,
        employee_id=employee.id,
        snapshot_id_number=employee.id_number,
        snapshot_employee_no=employee.employee_no,
        snapshot_employee_name=employee.name,
        snapshot_fixed_salary=Decimal("10000.00"),
        snapshot_performance_base=Decimal("2000.00"),
        gross_amount=Decimal("10000.00"),
        net_amount=Decimal("10000.00"),
    )
    db_session.add(record)
    db_session.add(
        CorrectionBatch(
            original_batch_id=original.id,
            replacement_batch_id=replacement.id,
            reason="修正导入数据",
        )
    )
    db_session.flush()

    persisted = db_session.scalar(select(PayrollRecord).where(PayrollRecord.id == record.id))
    assert persisted is not None
    assert persisted.snapshot_fixed_salary == Decimal("10000.00")


def test_locked_payroll_record_cannot_be_updated(db_session: Session) -> None:
    _, _, subject, period, employee = create_base_records(db_session)
    batch = PayrollBatch(
        subject_id=subject.id,
        payroll_period_id=period.id,
        batch_type="normal",
        status="draft",
    )
    db_session.add(batch)
    db_session.flush()
    record = PayrollRecord(
        payroll_batch_id=batch.id,
        payroll_period_id=period.id,
        subject_id=subject.id,
        employee_id=employee.id,
        snapshot_id_number=employee.id_number,
        snapshot_employee_no=employee.employee_no,
        snapshot_employee_name=employee.name,
        snapshot_fixed_salary=Decimal("10000.00"),
        snapshot_performance_base=Decimal("2000.00"),
    )
    db_session.add(record)
    db_session.flush()

    record.calculation_status = "locked"
    db_session.flush()
    record.snapshot_employee_name = "不应被修改"

    with pytest.raises(IntegrityError):
        db_session.flush()


def test_organization_and_employee_api_flow(api_client: TestClient, postgres_engine) -> None:
    company = api_client.post("/organization/company", json={"code": "ACME", "name": "示例公司"})
    assert company.status_code == 201
    company_id = company.json()["id"]
    assert (
        api_client.post(
            "/organization/company", json={"code": "OTHER", "name": "其他公司"}
        ).status_code
        == 409
    )
    assert api_client.patch("/organization/company", json={"name": "更新公司"}).status_code == 200
    assert api_client.get("/organization/company").json()["name"] == "更新公司"

    city = api_client.post("/organization/cities", json={"code": "BJ", "name": "北京"})
    subject = api_client.post(
        "/organization/subjects",
        json={"company_id": company_id, "code": "MAIN", "name": "主主体"},
    )
    department = api_client.post(
        "/organization/departments",
        json={"company_id": company_id, "code": "ENG", "name": "研发"},
    )
    subject_department = api_client.post(
        "/organization/subject-departments",
        json={
            "company_id": company_id,
            "subject_id": subject.json()["id"],
            "department_id": department.json()["id"],
            "code": "ENG",
            "name": "研发",
        },
    )
    assert all(
        response.status_code == 201 for response in (city, subject, department, subject_department)
    )
    for path, response in (
        ("cities", city),
        ("subjects", subject),
        ("departments", department),
        ("subject-departments", subject_department),
    ):
        row_id = response.json()["id"]
        assert any(row["id"] == row_id for row in api_client.get(f"/organization/{path}").json())
        assert (
            api_client.patch(f"/organization/{path}/{row_id}", json={"name": "已更新"}).status_code
            == 200
        )
        assert any(
            row["id"] == row_id and row["name"] == "已更新"
            for row in api_client.get(f"/organization/{path}").json()
        )

    duplicate_city = api_client.post("/organization/cities", json={"code": "BJ", "name": "重复"})
    assert duplicate_city.status_code == 409
    assert "城市编码" in duplicate_city.json()["detail"]
    self_parent = api_client.patch(
        f"/organization/departments/{department.json()['id']}",
        json={"parent_id": department.json()["id"]},
    )
    assert self_parent.status_code == 409
    assert "循环" in self_parent.json()["detail"]
    other_department = api_client.post(
        "/organization/departments",
        json={"company_id": company_id, "code": "OPS", "name": "运营"},
    )
    assert other_department.status_code == 201
    assert (
        api_client.patch(
            f"/organization/departments/{other_department.json()['id']}",
            json={"parent_id": department.json()["id"]},
        ).status_code
        == 200
    )
    cycle = api_client.patch(
        f"/organization/departments/{department.json()['id']}",
        json={"parent_id": other_department.json()["id"]},
    )
    assert cycle.status_code == 409
    assert "循环" in cycle.json()["detail"]
    cross_company = api_client.post(
        "/organization/subject-departments",
        json={
            "company_id": company_id + 999,
            "subject_id": subject.json()["id"],
            "department_id": department.json()["id"],
            "code": "OTHER",
            "name": "跨公司",
        },
    )
    assert cross_company.status_code == 400
    assert "同一公司" in cross_company.json()["detail"]

    employee_payload = {
        "company_id": company_id,
        "id_number": "11010119900101123X",
        "employee_no": "E001",
        "name": "测试员工",
        "employee_type": "employee",
        "formal_status": True,
        "probation_status": "confirmed",
        "hire_date": "2026-01-01",
        "assignment": {
            "subject_id": subject.json()["id"],
            "subject_department_id": subject_department.json()["id"],
            "position_title": "工程师",
            "effective_from": "2026-01-01",
        },
        "salary": {
            "fixed_salary": "8000.00",
            "performance_base": "2000.00",
            "effective_from": "2026-01-01",
        },
        "base": {"city_id": city.json()["id"], "effective_from": "2026-01-01"},
        "bank_account": {
            "account_number": "00006222000000000001",
            "account_name": "测试员工",
            "effective_from": "2026-01-01",
        },
    }
    created = api_client.post("/employees", json=employee_payload)
    assert created.status_code == 201
    employee_id = created.json()["id"]
    assert created.json()["bank_account"]["account_number"] == "00006222000000000001"
    immutable_id = api_client.patch(f"/employees/{employee_id}", json={"id_number": "不可修改"})
    assert immutable_id.status_code == 422
    assert (
        api_client.get(f"/employees/{employee_id}").json()["id_number"]
        == employee_payload["id_number"]
    )

    duplicate = api_client.post("/employees", json={**employee_payload, "employee_no": "E002"})
    assert duplicate.status_code == 409
    assert "身份证" in duplicate.json()["detail"]

    updated = api_client.patch(
        f"/employees/{employee_id}",
        json={
            "salary": {
                "fixed_salary": "8500.00",
                "performance_base": "2125.00",
                "effective_from": "2026-07-01",
            }
        },
    )
    assert updated.status_code == 200
    assert updated.json()["salary"]["fixed_salary"] == "8500.00"

    overlapping_salary = api_client.patch(
        f"/employees/{employee_id}",
        json={
            "salary": {
                "fixed_salary": "8500.00",
                "performance_base": "2125.00",
                "effective_from": "2026-07-01",
            }
        },
    )
    assert overlapping_salary.status_code == 400
    assert "生效日期" in overlapping_salary.json()["detail"]

    midmonth_raise = api_client.patch(
        f"/employees/{employee_id}",
        json={
            "salary": {
                "fixed_salary": "9000.00",
                "performance_base": "2250.00",
                "effective_from": "2026-08-15",
            }
        },
    )
    assert midmonth_raise.status_code == 400
    assert "工资期间的首日" in midmonth_raise.json()["detail"]

    failed_update = api_client.patch(
        f"/employees/{employee_id}",
        json={
            "name": "不得部分保存",
            "salary": {
                "fixed_salary": "9000.00",
                "performance_base": "2250.00",
                "effective_from": "2026-08-01",
            },
            "bank_account": {
                "account_number": "00006222000000000001",
                "account_name": "测试员工",
                "effective_from": "2026-08-01",
            },
        },
    )
    assert failed_update.status_code == 409
    unchanged = api_client.get(f"/employees/{employee_id}").json()
    assert unchanged["name"] == "测试员工"
    assert unchanged["salary"]["fixed_salary"] == "8500.00"

    with Session(postgres_engine) as session:
        period = PayrollPeriod(
            year=2026,
            month=7,
            period_start=date(2026, 7, 1),
            period_end=date(2026, 7, 31),
        )
        session.add(period)
        session.flush()
        batch = PayrollBatch(
            subject_id=subject.json()["id"],
            payroll_period_id=period.id,
            batch_type="normal",
            status="draft",
        )
        session.add(batch)
        session.flush()
        record = PayrollRecord(
            payroll_batch_id=batch.id,
            payroll_period_id=period.id,
            subject_id=subject.json()["id"],
            employee_id=employee_id,
            snapshot_id_number=employee_payload["id_number"],
            snapshot_employee_no=employee_payload["employee_no"],
            snapshot_employee_name="测试员工",
            snapshot_fixed_salary=Decimal("8500.00"),
            snapshot_performance_base=Decimal("2125.00"),
        )
        session.add(record)
        session.flush()
        record_id = record.id
        session.commit()

    deactivated = api_client.patch(f"/employees/{employee_id}", json={"active": False})
    assert deactivated.status_code == 200
    assert deactivated.json()["active"] is False
    persisted = api_client.get(f"/employees/{employee_id}")
    assert persisted.status_code == 200
    assert persisted.json()["salary"]["fixed_salary"] == "8500.00"
    with Session(postgres_engine) as session:
        historical = session.get(PayrollRecord, record_id)
        assert historical is not None
        assert historical.snapshot_employee_name == "测试员工"
        assert historical.snapshot_fixed_salary == Decimal("8500.00")

    with Session(postgres_engine) as session:
        salary_rows = list(
            session.scalars(
                select(EmployeeSalary)
                .where(EmployeeSalary.employee_id == employee_id)
                .order_by(EmployeeSalary.effective_from)
            )
        )
        assert salary_rows[0].effective_to == date(2026, 7, 1)

    blocked_delete = api_client.delete(f"/organization/cities/{city.json()['id']}")
    assert blocked_delete.status_code == 409
    assert "base 地" in blocked_delete.json()["detail"]

    probation_payload = {
        **employee_payload,
        "id_number": "11010119900101124X",
        "employee_no": "E003",
        "probation_status": "in_probation",
        "formal_status": False,
        "salary": {
            **employee_payload["salary"],
            "performance_base": "0.00",
        },
    }
    probation_employee = api_client.post("/employees", json=probation_payload)
    assert probation_employee.status_code == 201

    changed_fixed_salary = api_client.patch(
        f"/employees/{probation_employee.json()['id']}",
        json={
            "probation_status": "confirmed",
            "formal_status": True,
            "salary": {
                "fixed_salary": "8500.00",
                "performance_base": "2125.00",
                "effective_from": "2026-07-01",
            },
        },
    )
    assert changed_fixed_salary.status_code == 400
    assert "固定薪资必须保持不变" in changed_fixed_salary.json()["detail"]

    confirmed = api_client.patch(
        f"/employees/{probation_employee.json()['id']}",
        json={
            "probation_status": "confirmed",
            "formal_status": True,
            "salary": {
                "fixed_salary": "8000.00",
                "performance_base": "2000.00",
                "effective_from": "2026-07-01",
            },
        },
    )
    assert confirmed.status_code == 200
    assert confirmed.json()["probation_status"] == "confirmed"
    assert confirmed.json()["salary"]["fixed_salary"] == "8000.00"


def test_city_and_attendance_rules_api_flow(api_client: TestClient) -> None:
    city = api_client.post("/organization/cities", json={"code": "RULES-BJ", "name": "规则北京"})
    assert city.status_code == 201
    city_id = city.json()["id"]
    city_rule = api_client.post(
        "/rules/city-rules",
        json={
            "city_id": city_id,
            "effective_from": "2026-01-01",
            "effective_to": "2026-12-31",
            "version": "v2026.1",
            "source": "北京市政策文件",
            "fixed_base": "6821.00",
            "social_items": [
                {
                    "item_code": "pension",
                    "item_name": "养老",
                    "company_rate": "0.16",
                    "employee_rate": "0.08",
                }
            ],
        },
    )
    assert city_rule.status_code == 201
    assert city_rule.json()["fixed_base"] == "6821.00"
    assert city_rule.json()["housing_company_rate"] == "0.05000000"
    assert city_rule.json()["housing_base_source"] == "fixed_salary"
    assert len(city_rule.json()["social_items"]) == 1
    assert len(api_client.get("/rules/city-rules").json()) == 1

    overlapping = api_client.post(
        "/rules/city-rules",
        json={
            "city_id": city_id,
            "effective_from": "2026-06-01",
            "version": "v2026.2",
            "fixed_base": "7000.00",
            "social_items": [
                {
                    "item_code": "pension",
                    "item_name": "养老",
                    "company_rate": "0.16",
                    "employee_rate": "0.08",
                }
            ],
        },
    )
    assert overlapping.status_code == 409
    assert "重叠" in overlapping.json()["detail"]

    invalid_housing = api_client.post(
        "/rules/city-rules",
        json={
            "city_id": city_id,
            "effective_from": "2027-01-01",
            "version": "v2027.1",
            "fixed_base": "7000.00",
            "housing_company_rate": "0.06",
            "social_items": [
                {
                    "item_code": "pension",
                    "item_name": "养老",
                    "company_rate": "0.16",
                    "employee_rate": "0.08",
                }
            ],
        },
    )
    assert invalid_housing.status_code == 422
    assert "5%" in str(invalid_housing.json())

    attendance = api_client.post(
        "/rules/attendance",
        json={
            "effective_from": "2026-01-01",
            "effective_to": "2026-12-31",
            "version": "v2026.1",
            "source": "公司考勤制度",
        },
    )
    assert attendance.status_code == 201
    assert attendance.json()["standard_hours"] == "8.00"
    assert attendance.json()["missed_punch_amount"] == "30.00"
    assert attendance.json()["exempt_level_number"] == 7
    assert attendance.json()["makeup_punch_exempt"] is True
    assert len(api_client.get("/rules/attendance").json()) == 1

    overlapping_attendance = api_client.post(
        "/rules/attendance",
        json={"effective_from": "2026-06-01", "version": "v2026.2"},
    )
    assert overlapping_attendance.status_code == 409
    assert "重叠" in overlapping_attendance.json()["detail"]


def test_payroll_workbench_api_flow(api_client: TestClient) -> None:
    company = api_client.get("/organization/company")
    assert company.status_code == 200
    if company.json() is None:
        company = api_client.post(
            "/organization/company", json={"code": "PAY05", "name": "05 测试公司"}
        )
        assert company.status_code == 201
    company_id = company.json()["id"]
    subject = api_client.post(
        "/organization/subjects",
        json={"company_id": company_id, "code": "PAY05-MAIN", "name": "05 主体"},
    )
    assert subject.status_code == 201
    subject_id = subject.json()["id"]
    department = api_client.post(
        "/organization/departments",
        json={"company_id": company_id, "code": "PAY05-D", "name": "05 部门"},
    )
    assert department.status_code == 201
    subject_department = api_client.post(
        "/organization/subject-departments",
        json={
            "company_id": company_id,
            "subject_id": subject_id,
            "department_id": department.json()["id"],
            "code": "PAY05-SD",
            "name": "05 主体部门",
        },
    )
    assert subject_department.status_code == 201
    city = api_client.post("/organization/cities", json={"code": "PAY05-C", "name": "05 城市"})
    assert city.status_code == 201
    employee = api_client.post(
        "/employees",
        json={
            "company_id": company_id,
            "id_number": "11010119900101125X",
            "employee_no": "PAY05-E001",
            "name": "05 员工",
            "employee_type": "employee",
            "formal_status": True,
            "probation_status": "confirmed",
            "hire_date": "2020-01-01",
            "assignment": {
                "subject_id": subject_id,
                "subject_department_id": subject_department.json()["id"],
                "position_title": "工程师",
                "effective_from": "2020-01-01",
            },
            "salary": {
                "fixed_salary": "8000.00",
                "performance_base": "2000.00",
                "effective_from": "2020-01-01",
            },
            "base": {"city_id": city.json()["id"], "effective_from": "2020-01-01"},
            "bank_account": {
                "account_number": "6222000000000002",
                "account_name": "05 员工",
                "effective_from": "2020-01-01",
            },
        },
    )
    assert employee.status_code == 201, employee.text

    period = api_client.post("/payroll/periods", json={"period": "2026-02"})
    assert period.status_code == 200
    assert period.json()["period_start"] == "2026-02-01"
    assert period.json()["period_end"] == "2026-02-28"
    failed_batch = api_client.post(
        f"/payroll/periods/{period.json()['id']}/batches",
        json={"subject_id": 999999},
    )
    assert failed_batch.status_code == 404
    assert api_client.get("/payroll/workbench?period=2026-02").json()["batches"] == []
    batch = api_client.post(
        f"/payroll/periods/{period.json()['id']}/batches",
        json={"subject_id": subject_id},
    )
    assert batch.status_code == 201
    assert batch.json()["scope"]["employee_count"] == 1
    assert batch.json()["scope"]["source"] == "employee_assignment_for_period"
    assert batch.json()["data_preparation"]["attendance"]["status"] == "missing"

    duplicate = api_client.post(
        f"/payroll/periods/{period.json()['id']}/batches",
        json={"subject_id": subject_id, "batch_type": "normal"},
    )
    assert duplicate.status_code == 409
    assert "已有正常批次" in duplicate.json()["detail"]
    assert api_client.get("/payroll/workbench?period=2026-02").status_code == 200
    assert api_client.get("/payroll/workbench?period=2026-03").status_code == 404


def test_concurrent_normal_batch_creation_keeps_one_batch(
    db_session: Session, postgres_engine
) -> None:
    _, _, subject, period, _ = create_base_records(db_session, suffix="CONCURRENT")
    subject_id, period_id = subject.id, period.id
    db_session.commit()

    def create() -> str:
        with Session(postgres_engine) as session:
            try:
                create_batch(session, period_id, subject_id, "normal", None)
                return "created"
            except PayrollWorkbenchError as exc:
                return str(exc.status_code)

    with ThreadPoolExecutor(max_workers=2) as executor:
        results = list(executor.map(lambda _: create(), range(2)))

    assert sorted(results) == ["409", "created"]
    with Session(postgres_engine) as session:
        assert session.scalar(
            select(PayrollBatch.id).where(
                PayrollBatch.payroll_period_id == period_id,
                PayrollBatch.subject_id == subject_id,
                PayrollBatch.batch_type == "normal",
            )
        )


def test_employee_import_uploads_partial_rows_and_corrects_one_row(
    api_client, db_session: Session
) -> None:
    company = Company(code="ACMEIMPORT", name="Acme")
    city = City(code="BJIMPORT", name="北京")
    subject = Subject(company_id=1, code="MAINIMPORT", name="主主体")
    existing = Employee(
        company_id=1,
        id_number="11010119900101129X",
        employee_no="EXISTING",
        name="既有员工",
        employee_type="employee",
        formal_status=True,
        probation_status="confirmed",
        hire_date=date(2020, 1, 1),
    )
    db_session.add_all([company, city])
    db_session.flush()
    subject.company_id = company.id
    existing.company_id = company.id
    db_session.add_all([subject, existing])
    db_session.flush()
    department = Department(company_id=company.id, code="DEVIMPORT", name="研发")
    db_session.add(department)
    db_session.flush()
    db_session.add(
        SubjectDepartment(
            company_id=company.id,
            subject_id=subject.id,
            department_id=department.id,
            code="DEVIMPORT",
            name="研发",
        )
    )
    db_session.commit()

    def row(
        id_number: object, employee_no: str, fixed: str = "8000.00", performance: str = "2000.00"
    ) -> list[object]:
        return [
            id_number,
            employee_no,
            "导入员工",
            "employee",
            "2026-01-01",
            "",
            "confirmed",
            "2026-07-01",
            "是",
            subject.code,
            "DEVIMPORT",
            "工程师",
            "P6",
            "6",
            fixed,
            performance,
            city.code,
            "000012345678901234",
            "导入员工",
            "银行",
            "支行",
            "2026-01-01",
        ]

    workbook = load_workbook(BytesIO(make_template()))
    sheet = workbook["员工资料"]
    sheet.append(row("11010119900101123X", "IMP001"))
    sheet.append(row("11010119900101124X", "IMP002", performance="1000.00"))
    sheet.append(row("11010119900101123X", "IMP003"))
    sheet.append(row(110101199001011250, "IMP004"))
    content = BytesIO()
    workbook.save(content)

    response = api_client.post(
        f"/imports/employee-master?company_id={company.id}",
        files={
            "file": (
                "employees.xlsx",
                content.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert response.status_code == 201
    batch = response.json()
    assert (batch["total_rows"], batch["success_rows"], batch["error_rows"]) == (4, 1, 3)
    assert batch["template_version"] == "TPL-EMP-v2.4"
    assert batch["rows"][1]["errors"][0]["code"] == "SALARY_RATIO"
    assert any(error["code"] == "EMPLOYEE_EXISTS" for error in batch["rows"][2]["errors"])
    assert any(error["code"] == "IDENTIFIER_NOT_TEXT" for error in batch["rows"][3]["errors"])

    correction = api_client.post(
        f"/imports/{batch['id']}/rows/{batch['rows'][1]['id']}/correct",
        json={"values": {"绩效基数": "2000.00"}},
    )
    assert correction.status_code == 200
    assert correction.json()["success_rows"] == 2
    assert correction.json()["rows"][1]["correction_history"]
    assert db_session.scalar(select(Employee).where(Employee.employee_no == "IMP002"))
    assert existing.id


def test_attendance_import_partial_correction_and_revision(api_client, db_session: Session) -> None:
    company = Company(code="ATT06", name="考勤测试公司")
    city = City(code="ATT06CITY", name="考勤测试城市")
    db_session.add_all([company, city])
    db_session.flush()
    subject = Subject(company_id=company.id, code="ATT06SUBJECT", name="考勤测试主体")
    department = Department(company_id=company.id, code="ATT06DEP", name="考勤测试部门")
    period = PayrollPeriod(
        year=2027, month=2, period_start=date(2027, 2, 1), period_end=date(2027, 2, 28)
    )
    db_session.add_all([subject, department, period])
    db_session.flush()
    relation = SubjectDepartment(
        company_id=company.id,
        subject_id=subject.id,
        department_id=department.id,
        code="ATT06REL",
        name="考勤测试部门",
    )
    batch = PayrollBatch(subject_id=subject.id, payroll_period_id=period.id, batch_type="normal")
    db_session.add_all([relation, batch])
    db_session.flush()
    identities = ["11010119900101123X", "11010119900101124X", "11010119900101125X"]
    employees = []
    for index, identity in enumerate(identities):
        employee = Employee(
            company_id=company.id,
            id_number=identity,
            employee_no=f"ATT06-{index}",
            name=f"考勤员工{index}",
            employee_type="employee",
            hire_date=date(2020, 1, 1),
        )
        db_session.add(employee)
        db_session.flush()
        db_session.add(
            EmployeeAssignment(
                employee_id=employee.id,
                subject_id=subject.id,
                subject_department_id=relation.id,
                position_title="工程师",
                effective_from=date(2020, 1, 1),
            )
        )
        employees.append(employee)
    db_session.commit()

    def row(index: int, days: object = "20", late: object = 10, missed: object = 2) -> list[object]:
        return [
            identities[index],
            f"考勤员工{index}",
            days,
            late,
            5,
            "1.5",
            "0.5",
            missed,
            1,
            "考勤系统导出",
        ]

    workbook = load_workbook(BytesIO(make_attendance_template()))
    sheet = workbook["月度考勤"]
    sheet.append(row(0))
    sheet.append(row(1, days=0))
    sheet.append(row(0))
    sheet.append(row(2, late="=1+1"))
    content = BytesIO()
    workbook.save(content)
    files = {
        "file": (
            "attendance.xlsx",
            content.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    }
    response = api_client.post(f"/imports/attendance?payroll_batch_id={batch.id}", files=files)
    assert response.status_code == 201, response.text
    imported = response.json()
    assert (imported["success_rows"], imported["error_rows"]) == (1, 3)
    assert imported["rows"][1]["errors"][0]["code"] == "WORK_DAYS_OUT_OF_RANGE"
    assert any(item["code"] == "DUPLICATE_IN_FILE" for item in imported["rows"][2]["errors"])
    assert any(item["code"] == "FORMULA_NOT_ALLOWED" for item in imported["rows"][3]["errors"])
    db_session.expire_all()
    assert db_session.get(PayrollPeriod, period.id).attendance_input_revision == 1
    fact = db_session.scalar(
        select(AttendanceRecord).where(AttendanceRecord.employee_id == employees[0].id)
    )
    assert (fact.paid_leave_days, fact.unpaid_leave_days, fact.corrected_punch_count) == (
        Decimal("1.50"),
        Decimal("0.50"),
        1,
    )
    assert fact.leave_type == "mixed"
    assert (
        api_client.post(f"/imports/attendance?payroll_batch_id={batch.id}", files=files).status_code
        == 409
    )
    assert (
        api_client.get(f"/imports/attendance/{imported['id']}/file").content == content.getvalue()
    )

    correction = api_client.post(
        f"/imports/attendance/{imported['id']}/rows/{imported['rows'][1]['id']}/correct",
        json={"values": {"应出勤天数": "20"}},
    )
    assert correction.status_code == 200, correction.text
    assert correction.json()["success_rows"] == 2
    assert correction.json()["rows"][1]["correction_history"]
    db_session.expire_all()
    assert db_session.get(PayrollPeriod, period.id).attendance_input_revision == 2
    assert (
        api_client.post(
            f"/imports/attendance/{imported['id']}/rows/{imported['rows'][1]['id']}/correct",
            json={"values": {"应出勤天数": "20"}},
        ).status_code
        == 409
    )
    db_session.expire_all()
    assert db_session.get(PayrollPeriod, period.id).attendance_input_revision == 2

    second = load_workbook(BytesIO(make_attendance_template()))
    second["月度考勤"].append(row(0, late=12))
    second["月度考勤"].append(row(2, days="-1"))
    second["月度考勤"].append(row(2, days="29"))
    second["月度考勤"].append(row(2, late=20000))
    second["月度考勤"].append(row(2, days="1", missed=3))
    second_content = BytesIO()
    second.save(second_content)
    duplicate = api_client.post(
        f"/imports/attendance?payroll_batch_id={batch.id}",
        files={
            "file": (
                "again.xlsx",
                second_content.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert duplicate.status_code == 201
    assert duplicate.json()["success_rows"] == 0
    assert duplicate.json()["rows"][0]["errors"][0]["code"] == "DUPLICATE_PERIOD"
    assert any(error["code"] == "INVALID_NUMBER" for error in duplicate.json()["rows"][1]["errors"])
    assert any(
        error["code"] == "WORK_DAYS_OUT_OF_RANGE" for error in duplicate.json()["rows"][2]["errors"]
    )
    assert any(
        error["code"] == "MINUTES_OUT_OF_RANGE" for error in duplicate.json()["rows"][3]["errors"]
    )
    assert any(
        error["code"] == "PUNCH_OUT_OF_RANGE" for error in duplicate.json()["rows"][4]["errors"]
    )


def test_performance_import_partial_correction_and_probation(
    api_client, db_session: Session
) -> None:
    company = Company(code="PERF07", name="绩效测试公司")
    db_session.add(company)
    db_session.flush()
    subject = Subject(company_id=company.id, code="PERF07SUB", name="绩效测试主体")
    department = Department(company_id=company.id, code="PERF07DEP", name="绩效测试部门")
    period = PayrollPeriod(
        year=2027, month=3, period_start=date(2027, 3, 1), period_end=date(2027, 3, 31)
    )
    db_session.add_all([subject, department, period])
    db_session.flush()
    relation = SubjectDepartment(
        company_id=company.id,
        subject_id=subject.id,
        department_id=department.id,
        code="PERF07REL",
        name="绩效测试部门",
    )
    batch = PayrollBatch(subject_id=subject.id, payroll_period_id=period.id, batch_type="normal")
    db_session.add_all([relation, batch])
    db_session.flush()
    identities = [f"1101011990010112{i}X" for i in range(7)]
    employees = []
    for index, identity in enumerate(identities):
        employee = Employee(
            company_id=company.id,
            id_number=identity,
            employee_no=f"PERF07-{index}",
            name=f"绩效员工{index}",
            employee_type="employee",
            hire_date=date(2020, 1, 1),
            probation_status="in_probation" if index == 4 else "confirmed",
            probation_date=date(2027, 4, 15)
            if index == 4
            else (date(2027, 3, 15) if index == 5 else None),
        )
        db_session.add(employee)
        db_session.flush()
        db_session.add(
            EmployeeAssignment(
                employee_id=employee.id,
                subject_id=subject.id,
                subject_department_id=relation.id,
                position_title="工程师",
                effective_from=date(2020, 1, 1),
            )
        )
        employees.append(employee)
    db_session.commit()

    preparation = api_client.get("/payroll/workbench?period=2027-03").json()["batches"][0][
        "data_preparation"
    ]["performance"]
    assert preparation["missing_count"] == 6  # 试用期员工不需要绩效记录

    workbook = load_workbook(BytesIO(make_performance_template()))
    sheet = workbook["月度绩效"]
    for index, value in enumerate(["0", "", "123456789012345.123456789", "-1", "1", "1.2", "=1+1"]):
        sheet.append([identities[index], f"绩效员工{index}", value, "绩效表"])
    sheet.append([identities[0], "绩效员工0", "1", "重复行"])
    content = BytesIO()
    workbook.save(content)
    files = {
        "file": (
            "performance.xlsx",
            content.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    }
    response = api_client.post(f"/imports/performance?payroll_batch_id={batch.id}", files=files)
    assert response.status_code == 201, response.text
    imported = response.json()
    assert (imported["success_rows"], imported["error_rows"]) == (3, 5)
    assert imported["rows"][0]["normalized_data"]["coefficient"] == "0"
    assert imported["rows"][1]["errors"][0]["code"] == "COEFFICIENT_REQUIRED"
    assert imported["rows"][3]["errors"][0]["code"] == "INVALID_COEFFICIENT"
    assert any(e["code"] == "PROBATION_NO_PERFORMANCE" for e in imported["rows"][4]["errors"])
    assert imported["rows"][5]["normalized_data"]["coefficient"] == "1.2"
    assert any(e["code"] == "FORMULA_NOT_ALLOWED" for e in imported["rows"][6]["errors"])
    assert any(e["code"] == "DUPLICATE_IN_FILE" for e in imported["rows"][7]["errors"])
    db_session.expire_all()
    assert db_session.get(PayrollPeriod, period.id).performance_input_revision == 3
    high = db_session.scalar(
        select(PerformanceRecord).where(PerformanceRecord.employee_id == employees[2].id)
    )
    assert high.coefficient == Decimal("123456789012345.123456789")
    assert high.source_value == "123456789012345.123456789"
    assert (
        api_client.get(f"/imports/performance/{imported['id']}/file").content == content.getvalue()
    )
    assert (
        api_client.post(
            f"/imports/performance?payroll_batch_id={batch.id}", files=files
        ).status_code
        == 409
    )

    for index, value in [(1, "0.8"), (3, "1.5"), (6, "1")]:
        correction = api_client.post(
            f"/imports/performance/{imported['id']}/rows/{imported['rows'][index]['id']}/correct",
            json={"values": {"绩效系数": value}},
        )
        assert correction.status_code == 200, correction.text
        assert correction.json()["rows"][index]["correction_history"]
    db_session.expire_all()
    assert db_session.get(PayrollPeriod, period.id).performance_input_revision == 6
    assert (
        api_client.post(
            f"/imports/performance/{imported['id']}/rows/{imported['rows'][1]['id']}/correct",
            json={"values": {"绩效系数": "0.8"}},
        ).status_code
        == 409
    )
    assert (
        api_client.get("/payroll/workbench?period=2027-03").json()["batches"][0][
            "data_preparation"
        ]["performance"]["status"]
        == "ready"
    )

    another = load_workbook(BytesIO(make_performance_template()))
    another["月度绩效"].append([identities[0], "绩效员工0", "2", "再次导入"])
    second_content = BytesIO()
    another.save(second_content)
    duplicate = api_client.post(
        f"/imports/performance?payroll_batch_id={batch.id}",
        files={"file": ("again.xlsx", second_content.getvalue(), files["file"][2])},
    )
    assert duplicate.status_code == 201
    assert duplicate.json()["rows"][0]["errors"][0]["code"] == "DUPLICATE_PERIOD"

    # Downgrade cannot narrow this large coefficient; clear its test fact.
    db_session.delete(high)
    db_session.commit()


def test_ordinary_trial_is_partial_versioned_and_outside_ledger(
    api_client, db_session: Session, monkeypatch
) -> None:
    company = Company(code="TRIAL08", name="试算测试公司")
    city = City(code="TRIAL08-C", name="测试城市")
    missing_city = City(code="TRIAL08-M", name="缺规则城市")
    db_session.add_all([company, city, missing_city])
    db_session.flush()
    subject = Subject(company_id=company.id, code="TRIAL08-S", name="测试主体")
    department = Department(company_id=company.id, code="TRIAL08-D", name="测试部门")
    period = PayrollPeriod(
        year=2028, month=3, period_start=date(2028, 3, 1), period_end=date(2028, 3, 31)
    )
    db_session.add_all([subject, department, period])
    db_session.flush()
    relation = SubjectDepartment(
        company_id=company.id,
        subject_id=subject.id,
        department_id=department.id,
        code="TRIAL08-R",
        name="测试部门",
    )
    batch = PayrollBatch(subject_id=subject.id, payroll_period_id=period.id, batch_type="normal")
    city_rule = SocialSecurityRule(
        city_id=city.id,
        effective_from=date(2028, 1, 1),
        version="trial-08",
        fixed_base=Decimal("5000"),
    )
    housing = HousingFundRule(
        city_id=city.id,
        effective_from=date(2028, 1, 1),
        company_rate=Decimal("0.05"),
        employee_rate=Decimal("0.05"),
        base_min=0,
        base_max=0,
        version="trial-08",
        base_source="fixed_salary",
    )
    attendance_rule = AttendanceRule(
        effective_from=date(2028, 1, 1),
        version="trial-08",
        standard_hours=8,
        missed_punch_amount=30,
        exempt_level_number=7,
        makeup_punch_exempt=True,
    )
    db_session.add_all([relation, batch, city_rule, housing, attendance_rule])
    db_session.flush()
    db_session.add(
        SocialSecurityItemRule(
            social_security_rule_id=city_rule.id,
            item_code="pension",
            item_name="养老",
            employee_rate=Decimal("0.10"),
            company_rate=Decimal("0.20"),
            base_min=0,
            base_max=0,
        )
    )
    employees = []
    for index in range(3):
        employee = Employee(
            company_id=company.id,
            id_number=f"1101011990010112{index}X",
            employee_no=f"TRIAL08-{index}",
            name=f"试算员工{index}",
            employee_type="employee",
            level_number=6,
            probation_status="in_probation" if index == 1 else "confirmed",
            hire_date=date(2020, 1, 1),
        )
        db_session.add(employee)
        db_session.flush()
        db_session.add_all(
            [
                EmployeeAssignment(
                    employee_id=employee.id,
                    subject_id=subject.id,
                    subject_department_id=relation.id,
                    position_title="工程师",
                    level_number=6,
                    effective_from=date(2020, 1, 1),
                ),
                EmployeeSalary(
                    employee_id=employee.id,
                    fixed_salary=Decimal("8000"),
                    performance_base=Decimal("0" if index == 1 else "2000"),
                    effective_from=date(2020, 1, 1),
                ),
                EmployeeBase(
                    employee_id=employee.id,
                    city_id=missing_city.id if index == 2 else city.id,
                    effective_from=date(2020, 1, 1),
                ),
            ]
        )
        employees.append(employee)
    db_session.flush()
    att_import = ImportBatch(
        company_id=company.id,
        payroll_period_id=period.id,
        payroll_batch_id=batch.id,
        import_type="attendance",
        original_filename="attendance.xlsx",
        file_sha256="8" * 64,
        template_version="TPL-ATT-v1",
        status="imported",
    )
    perf_import = ImportBatch(
        company_id=company.id,
        payroll_period_id=period.id,
        payroll_batch_id=batch.id,
        import_type="performance",
        original_filename="performance.xlsx",
        file_sha256="9" * 64,
        template_version="TPL-PERF-v1",
        status="imported",
    )
    db_session.add_all([att_import, perf_import])
    db_session.flush()
    for index, employee in enumerate(employees):
        db_session.add(
            AttendanceRecord(
                import_batch_id=att_import.id,
                payroll_period_id=period.id,
                employee_id=employee.id,
                expected_work_days=Decimal("20"),
                late_minutes=60 if index == 0 else 0,
                missed_punch_count=1 if index == 0 else 0,
            )
        )
        if index != 1:
            db_session.add(
                PerformanceRecord(
                    import_batch_id=perf_import.id,
                    payroll_period_id=period.id,
                    employee_id=employee.id,
                    coefficient=Decimal("1"),
                    source_value="1",
                )
            )
    db_session.commit()

    batch_response = api_client.get(f"/payroll/batches/{batch.id}")
    assert batch_response.status_code == 200
    assert batch_response.json()["scope"]["employee_count"] == 3
    response = api_client.post(f"/payroll/batches/{batch.id}/trial")
    assert response.status_code == 200, response.text
    trial = response.json()
    assert (trial["success_count"], trial["error_count"]) == (2, 1)
    assert trial["includes_final_incentive"] is False
    assert trial["ready_for_confirmation"] is False
    assert trial["results"][0]["amounts"]["untaxed_amount"] == "9020.00"
    assert trial["results"][1]["amounts"]["performance"] == "0.00"
    assert trial["results"][2]["errors"][0]["code"] == "CITY_RULE_MISSING"
    assert trial["results"][0]["steps"]
    assert (
        db_session.scalar(
            select(PayrollRecord.id).where(PayrollRecord.payroll_batch_id == batch.id)
        )
        is None
    )
    assert api_client.post(f"/payroll/batches/{batch.id}/trial").json()["id"] == trial["id"]

    db_session.expire_all()
    salary = db_session.scalar(
        select(EmployeeSalary).where(EmployeeSalary.employee_id == employees[0].id)
    )
    salary.fixed_salary = Decimal("8800")
    salary.performance_base = Decimal("2200")
    db_session.commit()
    stale = api_client.get(f"/payroll/batches/{batch.id}/trial").json()
    assert stale["stale"] is True
    assert "试算输入已变化，请重新试算" in stale["confirmation_blockers"]
    with monkeypatch.context() as patch:
        patch.setattr(
            payroll_trial_service,
            "calculate",
            lambda _value: (_ for _ in ()).throw(RuntimeError("calculation aborted")),
        )
        with pytest.raises(RuntimeError, match="calculation aborted"):
            api_client.post(f"/payroll/batches/{batch.id}/trial")
    db_session.expire_all()
    assert len(list(db_session.scalars(select(PayrollTrialRun)))) == 1
    rerun = api_client.post(f"/payroll/batches/{batch.id}/trial")
    assert rerun.status_code == 200, rerun.text
    assert rerun.json()["id"] != trial["id"]
    assert rerun.json()["results"][0]["amounts"]["untaxed_amount"] == "9975.00"
    assert db_session.scalar(
        select(PayrollTrialRun.id).where(PayrollTrialRun.payroll_batch_id == batch.id)
    )
    probation_employee = db_session.get(Employee, employees[1].id)
    probation_employee.probation_status = "confirmed"
    probation_employee.probation_date = date(2028, 3, 16)
    probation_salary = db_session.scalar(
        select(EmployeeSalary).where(EmployeeSalary.employee_id == probation_employee.id)
    )
    probation_salary.effective_to = date(2028, 3, 16)
    db_session.add_all(
        [
            EmployeeSalary(
                employee_id=probation_employee.id,
                fixed_salary=Decimal("8000"),
                performance_base=Decimal("2000"),
                effective_from=date(2028, 3, 16),
            ),
            PerformanceRecord(
                import_batch_id=perf_import.id,
                payroll_period_id=period.id,
                employee_id=probation_employee.id,
                coefficient=Decimal("1"),
                source_value="1",
            ),
        ]
    )
    db_session.commit()
    promoted = api_client.post(f"/payroll/batches/{batch.id}/trial").json()
    assert promoted["results"][1]["amounts"]["fixed"] == "8000.00"
    assert promoted["results"][1]["amounts"]["performance"] == "2000.00"

    item_rule = db_session.scalar(
        select(SocialSecurityItemRule).where(
            SocialSecurityItemRule.social_security_rule_id == city_rule.id
        )
    )
    item_rule.employee_rate = Decimal("0.11")
    db_session.commit()
    assert api_client.get(f"/payroll/batches/{batch.id}/trial").json()["stale"] is True

    other_subject = Subject(company_id=company.id, code="TRIAL08-S2", name="另一主体")
    db_session.add(other_subject)
    db_session.flush()
    other_relation = SubjectDepartment(
        company_id=company.id,
        subject_id=other_subject.id,
        department_id=department.id,
        code="TRIAL08-R2",
        name="另一主体部门",
    )
    db_session.add(other_relation)
    db_session.flush()
    assignment = db_session.scalar(
        select(EmployeeAssignment).where(EmployeeAssignment.employee_id == employees[0].id)
    )
    assignment.effective_to = date(2028, 3, 16)
    db_session.add(
        EmployeeAssignment(
            employee_id=employees[0].id,
            subject_id=other_subject.id,
            subject_department_id=other_relation.id,
            position_title="工程师",
            level_number=6,
            effective_from=date(2028, 3, 16),
        )
    )
    db_session.commit()
    moved = api_client.post(f"/payroll/batches/{batch.id}/trial").json()
    assert any(
        error["code"] == "CROSS_SUBJECT_UNSUPPORTED" for error in moved["results"][0]["errors"]
    )
    assert moved["results"][0]["snapshot"]["last_effective_subject_id"] == other_subject.id
