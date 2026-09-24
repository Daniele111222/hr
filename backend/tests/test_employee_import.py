from io import BytesIO

from openpyxl import load_workbook

from paylite.excel.employee_template import (
    HEADERS,
    SHEET_NAME,
    make_template,
    parse_workbook,
)


def test_employee_template_round_trip_and_preserves_text_identifiers() -> None:
    workbook = load_workbook(BytesIO(make_template()))
    sheet = workbook[SHEET_NAME]
    sheet.append(
        [
            "11010119900101123X",
            "0007",
            "张三",
            "employee",
            "2026-01-01",
            "",
            "confirmed",
            "2026-07-01",
            "是",
            "MAIN",
            "DEV",
            "工程师",
            "P6",
            "6",
            "8000.00",
            "2000.00",
            "BJ",
            "000012345678901234",
            "张三",
            "银行",
            "支行",
            "2026-01-01",
        ]
    )
    output = BytesIO()
    workbook.save(output)
    rows = parse_workbook(output.getvalue())

    assert rows[0].number == 3
    assert rows[0].raw["身份证号"] == "11010119900101123X"
    assert rows[0].raw["员工编号"] == "0007"
    assert rows[0].raw["银行卡号"] == "000012345678901234"
    assert rows[0].errors == []


def test_employee_import_records_formula_and_numeric_identifier_errors() -> None:
    workbook = load_workbook(BytesIO(make_template()))
    sheet = workbook[SHEET_NAME]
    sheet.append(["11010119900101123X"] + [None] * (len(HEADERS) - 1))
    sheet.cell(3, 2).value = 123456789012345678
    sheet.cell(3, 15).value = "=1+1"
    output = BytesIO()
    workbook.save(output)
    rows = parse_workbook(output.getvalue())

    codes = {error["code"] for error in rows[0].errors}
    assert "IDENTIFIER_NOT_TEXT" in codes
    assert "FORMULA_NOT_ALLOWED" in codes
