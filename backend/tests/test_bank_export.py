from io import BytesIO

from openpyxl import load_workbook


def test_bank_template_maps_all_sheets_expands_and_preserves_identifiers():
    from paylite.excel.bank_export import render_bank

    sheets = [
        "导入模版（北分）",
        "导入模版（成都）",
        "导入模版（杭州）",
        "导入模版（深圳）",
        "导入模版（武汉）",
    ]
    rows = [
        dict(
            subject_id=i,
            name="=测试",
            id_number="001101199001011234",
            bank_account="006222000000000000001",
            amount="900.10",
            remark="测试主体；未扣个税",
        )
        for i in range(1, 6)
        for _ in range(60)
    ]
    book = load_workbook(BytesIO(render_bank(rows, {str(i): s for i, s in enumerate(sheets, 1)})))
    assert book.sheetnames == sheets
    for sheet in book:
        wuhan = sheet.title == "导入模版（武汉）"
        account, name, amount = ("C", "E", "F") if wuhan else ("B", "D", "E")
        assert sheet[f"{account}2"].value == "006222000000000000001"
        assert sheet[f"{account}2"].data_type == "s"
        assert sheet[f"{name}2"].data_type == "s"
        assert sheet[f"{amount}61"].value == 900.1
        assert sheet[f"{amount}62"].value == f"=SUM({amount}2:{amount}61)"
        assert sheet[f"{amount}1"].value == "未扣个税金额"
        assert sheet["D2" if wuhan else "C2"].value is None
        if wuhan:
            assert sheet["A2"].value == "001101199001011234"
            assert sheet["A2"].data_type == "s"
        assert sheet[f"{account}61"]._style == sheet[f"{account}2"]._style
    assert not book._external_links


def test_runtime_template_is_clean_and_same_layout_can_serve_two_subjects():
    from pathlib import Path

    from paylite.excel.bank_export import TEMPLATE_PATH, render_bank

    source = load_workbook(Path(__file__).parents[2] / "docs/代发工资模板.xlsx")
    runtime = load_workbook(TEMPLATE_PATH)
    assert runtime.sheetnames == source.sheetnames
    for sheet in runtime:
        assert str(sheet.merged_cells) == str(source[sheet.title].merged_cells)
        for row in sheet.iter_rows(min_row=2):
            assert all(c.value is None for c in row)
        assert sheet["B1"]._style == source[sheet.title]["B1"]._style
    row = dict(
        subject_id=1,
        name="测试",
        id_number="001101199001011234",
        bank_account="001234567890123456789",
        amount="0.00",
        remark="未扣个税",
    )
    book = load_workbook(
        BytesIO(
            render_bank(
                [row, {**row, "subject_id": 2, "amount": "20.00"}],
                {"1": "导入模版（武汉）", "2": "导入模版（武汉）"},
            )
        )
    )
    assert book["导入模版（武汉）"]["F2"].value == 0
    assert book["导入模版（武汉）-S2"]["F2"].value == 20
    assert book["导入模版（北分）"]["B2"].value is None


def test_bank_rejects_negative_or_conflicting_accounts_before_rendering():
    from paylite.services.payroll_export import prepare_bank_rows

    row = {
        "subject_id": 1,
        "subject_name": "测试",
        "employee_id": 1,
        "payroll_batch_id": 1,
        "batch_type": "normal",
        "batch_no": 1,
        "snapshot": {"name": "测试", "id_number": "001", "bank_account": "00123"},
        "amounts": {"untaxed_amount": "-1.00"},
    }
    ledger = {"period": "2026-09", "records": [row]}
    assert any("为负" in b for b in prepare_bank_rows(ledger)[1])
    row["amounts"]["untaxed_amount"] = "10.00"
    ledger["records"].append(
        {**row, "payroll_batch_id": 2, "snapshot": {**row["snapshot"], "bank_account": "00456"}}
    )
    assert any("不一致" in b for b in prepare_bank_rows(ledger)[1])
