from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from paylite.excel.payroll_export import render_payroll


def test_payroll_template_expands_and_preserves_text_and_unknown_tax():
    rows = [
        {
            "employee_id": n,
            "subject_id": 1,
            "subject_name": "测试主体",
            "batch_type": "normal",
            "batch_no": 1,
            "batch_name": None,
            "snapshot": {
                "name": "=测试",
                "id_number": "110101199001011234",
                "employee_no": str(n),
                "bank_account": "6222000000000000001",
                "fixed_salary": "800.00",
                "performance_base": "200.00",
            },
            "amounts": {"gross": "1000.00", "untaxed_amount": "900.00", "employer_cost": "1200.00"},
            "items": [
                {"code": "housing_employee", "amount": "40.00"},
                {"code": "social_employee_pension", "amount": "60.00"},
            ],
        }
        for n in range(20)
    ]
    content = render_payroll({"period": "2026-09", "records": rows})
    book = load_workbook(BytesIO(content))
    detail = book["工资表明细"]
    assert detail["D5"].value == "110101199001011234"
    assert detail["BN5"].value == "6222000000000000001"
    assert detail["C5"].data_type == "s"
    assert detail["AU5"].value is None
    assert detail["AX5"].value == 900
    assert detail["AX24"].value == 900
    assert detail["AX4"].value == "=SUBTOTAL(9,AX5:AX24)"
    assert sum(Decimal(str(detail[f"AX{n}"].value)) for n in range(5, 25)) == Decimal("18000")
    assert "A1:AC1" in str(detail.merged_cells)
    assert book["汇总"]["D3"].value is None
    assert book["个税计算"]["B5"].value is None
    assert book["汇总"]["C3"].value == "=SUM('工资表明细'!AX5:AX24)"


def test_subject_ids_separate_identical_names_and_template_notes_survive():
    row = {
        "employee_id": 1,
        "subject_id": 1,
        "subject_name": "测试*",
        "batch_type": "supplement",
        "batch_no": 1,
        "batch_name": "补发",
        "snapshot": {"name": "测试", "id_number": "110101199001011234"},
        "amounts": {"gross": "100.00", "untaxed_amount": "100.00", "employer_cost": "100.00"},
        "items": [],
    }
    book = load_workbook(
        BytesIO(
            render_payroll(
                {"period": "2026-09", "records": [row, {**row, "subject_id": 2, "employee_id": 2}]}
            )
        )
    )
    assert book["汇总"]["C3"].value == "=SUM('工资表明细'!AX5:AX5)"
    assert book["汇总"]["C4"].value == "=SUM('工资表明细'!AX6:AX6)"
    assert book["导入模版（中汽）"]["A12"].value == "输入项"
    assert book["导入模版（北分）"]["B2"].value == '=IF(A2="","",ROW()-1)'
    assert book["工资表明细"]["AT5"].value is None


def test_runtime_template_preserves_sheet_headers_merges_and_styles():
    from pathlib import Path

    from paylite.excel.payroll_export import TEMPLATE_PATH

    source = load_workbook(Path(__file__).parents[2] / "docs/工资表模板.xlsx")
    runtime = load_workbook(TEMPLATE_PATH)
    assert runtime.sheetnames == source.sheetnames
    for sheet in source:
        target = runtime[sheet.title]
        assert str(target.merged_cells) == str(sheet.merged_cells)
        assert target.sheet_state == sheet.sheet_state
    for coordinate in ("B3", "D3", "AO3", "BN3"):
        assert runtime["工资表明细"][coordinate].value == source["工资表明细"][coordinate].value
        assert runtime["工资表明细"][coordinate]._style == source["工资表明细"][coordinate]._style
    assert runtime["工资表明细"]["AX3"].value == "未扣个税金额"
    assert not runtime._external_links
