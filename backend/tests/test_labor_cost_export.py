from decimal import Decimal
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook


def test_labor_cost_separates_personal_deductions_and_company_cost_across_subjects():
    from paylite.excel.labor_cost_export import render_labor_cost

    row = {
        "subject_id": 1,
        "subject_name": "=测试主体*",
        "employee_id": 1,
        "payroll_batch_id": 1,
        "batch_type": "normal",
        "batch_no": 1,
        "snapshot": {
            "name": "=员工",
            "id_number": "001101199001011234",
            "bank_account": "006222000000000000001",
            "department_name": "测试部门",
        },
        "amounts": {"gross": "1000.00", "untaxed_amount": "900.00", "employer_cost": "1200.00"},
        "items": [
            {"code": "social_employee_pension", "amount": "60.00"},
            {"code": "housing_employee", "amount": "40.00"},
            {"code": "social_company_pension", "amount": "160.00"},
            {"code": "housing_company", "amount": "40.00"},
        ],
    }
    supplement = {
        **row,
        "payroll_batch_id": 2,
        "batch_type": "supplement",
        "batch_no": 2,
        "batch_name": "补发原因",
        "amounts": {"gross": "50.00", "untaxed_amount": "50.00", "employer_cost": "50.00"},
        "items": [{"code": "supplement", "amount": "50.00"}],
    }
    rows = [row, supplement, {**row, "subject_id": 2}]
    book = load_workbook(BytesIO(render_labor_cost({"period": "2026-09", "records": rows})))
    detail = book["工资明细"]
    assert detail["AX5"].value == 900
    assert detail["BF5"].value == 1200
    assert [detail[f"{c}5"].value for c in ("AS", "AT", "BD", "BE")] == [60, 40, 160, 40]
    assert [detail[f"{c}6"].value for c in ("AS", "AT", "BD", "BE")] == [0, 0, 0, 0]
    assert detail["BF4"].value == "=SUM(BF5:BF7)"
    assert sum(Decimal(str(detail[f"BF{n}"].value)) for n in (5, 6, 7)) == Decimal("2450.00")
    assert sum(Decimal(str(detail[f"AX{n}"].value)) for n in (5, 6, 7)) == Decimal("1850.00")
    assert detail["D5"].value == "001101199001011234"
    assert detail["BN5"].value == "006222000000000000001"
    assert detail["C5"].data_type == "s"
    assert "独立补发" in detail["BG6"].value
    assert "补发原因" in detail["BG6"].value
    first, second = book.worksheets[:2]
    assert first["F1"].value == second["F1"].value == "=测试主体*"
    assert first["F1"].data_type == "s"
    assert first["D6"].value == 1  # A supplement does not increase headcount.
    assert first["N6"].value == "=SUM('工资明细'!AX5:AX6)"
    assert second["N6"].value == "=SUM('工资明细'!AX7:AX7)"
    assert first["E2"].value == "=SUM('工资明细'!BF5:BF6)"
    assert first["E31"].value == "=SUM('工资明细'!BE5:BE6)"
    assert first["F31"].value == "=SUM('工资明细'!AT5:AT6)"
    assert first["K56"].value == "=SUM('工资明细'!AS5:AS6)"
    assert first["E26"].value == "=SUM(E6:E25)"
    for s in book.worksheets[:-1]:
        assert s["K6"].value is None  # Unknown tax stays blank at every level.
        for c in s._cells.values():
            assert c.data_type != "e"
            assert "#REF!" not in str(c.value)
    assert detail["AU4"].value is None
    assert detail["AU5"].value is None
    assert detail["AX3"].value == "未扣个税金额"


def test_labor_cost_preserves_template_structure_and_expands_without_sample_data():
    from paylite.excel.labor_cost_export import TEMPLATE_PATH, render_labor_cost

    source = load_workbook(Path(__file__).parents[2] / "docs/人工成本表-模版.xlsx")
    runtime = load_workbook(TEMPLATE_PATH)
    assert runtime.sheetnames == source.sheetnames
    assert not runtime._external_links
    assert not runtime.defined_names
    for s in source:
        assert str(runtime[s.title].merged_cells) == str(s.merged_cells)
        assert runtime[s.title].sheet_state == s.sheet_state
        assert runtime[s.title]["A4"]._style == s["A4"]._style
    for s in runtime.worksheets[:-1]:
        assert s["F1"].value is None
        assert s["A6"].value is None
        assert s["K6"].value is None
    row = {
        "subject_id": 1,
        "subject_name": "主体",
        "employee_id": 1,
        "payroll_batch_id": 1,
        "batch_type": "supplement",
        "batch_no": 1,
        "batch_name": "补发",
        "snapshot": {"name": "员工", "id_number": "001"},
        "amounts": {"gross": "0.01", "untaxed_amount": "0.01", "employer_cost": "0.01"},
        "items": [],
    }
    rows = [{**row, "subject_id": n, "employee_id": i} for n in range(1, 9) for i in range(60)]
    book = load_workbook(BytesIO(render_labor_cost({"period": "2026-09", "records": rows})))
    assert book.sheetnames[:8] == runtime.sheetnames
    assert len(book.sheetnames) == 9
    detail = book["工资明细"]
    assert detail["AX484"].value == 0.01
    assert detail["AX4"].value == "=SUM(AX5:AX484)"
    assert detail["AX484"]._style == detail["AX5"]._style
    assert sum(Decimal(str(detail[f"BF{n}"].value)) for n in range(5, 485)) == Decimal("4.80")
    assert book.worksheets[-1]["D6"].value == 60
    assert book.worksheets[1].sheet_state == "visible"
