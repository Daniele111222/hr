"""清除人工成本历史样例和外部引用，保留工作表、样式与有效合计公式。"""

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell

root = Path(__file__).resolve().parents[2]
book = load_workbook(root / "docs/人工成本表-模版.xlsx", keep_links=False)
for sheet in book:
    if sheet.title == "工资明细":
        for cell in list(sheet._cells.values()):
            if not isinstance(cell, MergedCell) and cell.row > 3:
                cell.value = None
        sheet["A1"] = "未扣个税金额；无来源税务、项目工时字段留空"
        sheet["AX2"] = sheet["AX3"] = "未扣个税金额"
    else:
        totals = [r for r in range(1, sheet.max_row + 1) if sheet.cell(r, 1).value == "合计"]
        headers = [r for r in range(1, sheet.max_row + 1) if sheet.cell(r, 1).value == "一级部门"]
        for header, total in zip(headers, totals, strict=True):
            for row in sheet.iter_rows(min_row=header + 2, max_row=total - 1):
                for cell in row:
                    if not isinstance(cell, MergedCell):
                        cell.value = None
        # Unknown tax and after-tax adjustments must not SUM empty cells into zero.
        for row in sheet.iter_rows(
            min_row=headers[0] + 2, max_row=totals[0], min_col=11, max_col=13
        ):
            for cell in row:
                cell.value = None
        sheet["B1"] = sheet["F1"] = None
        sheet["N4"] = "未扣个税金额"
    for cell in list(sheet._cells.values()):
        if isinstance(cell, MergedCell):
            continue
        cell.comment = cell.hyperlink = None
        if cell.data_type == "e" or (
            cell.data_type == "f" and ("[" in cell.value or "#REF!" in cell.value)
        ):
            cell.value = None
    sheet.oddFooter.center.text = "未扣个税金额；项目工时无来源留空，公式需 Excel 打开后重算"
    sheet.auto_filter.ref = None
    sheet.print_area = None
book.defined_names.clear()
book.save(root / "backend/src/paylite/excel/templates/labor-cost-v1.xlsx")
