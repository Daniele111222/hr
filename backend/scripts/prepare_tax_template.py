"""用 LibreOffice 真正转换旧式模板；原件不变，运行时不依赖转换工具。"""

import subprocess
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

root = Path(__file__).resolve().parents[2]
source = root / "docs/正常工资薪金所得_模板.xls"
with TemporaryDirectory() as directory:
    temporary = Path(directory)
    subprocess.run(
        [
            "soffice",
            "--headless",
            f"-env:UserInstallation={(temporary / 'profile').as_uri()}",
            "--convert-to",
            "xlsx",
            "--outdir",
            directory,
            str(source),
        ],
        check=True,
    )
    book = load_workbook(temporary / source.with_suffix(".xlsx").name)
    assert book.sheetnames == ["Sheet1"]
    assert book.active.max_row == 1 and book.active.max_column == 33
    assert not book._external_links and not book.defined_names
    assert all(c.data_type not in {"f", "e"} for row in book.active for c in row)
    book.save(root / "backend/src/paylite/excel/templates/tax-v1.xlsx")
