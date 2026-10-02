"""按有效锁定台账写员工明细与主体成本，不重新算薪或分摊项目工时。"""

from copy import copy
from decimal import Decimal
from hashlib import sha256
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from paylite.excel.payroll_export import _write

TEMPLATE_PATH = Path(__file__).with_name("templates") / "labor-cost-v1.xlsx"
TEMPLATE_VERSION = "labor-cost-v1-" + sha256(TEMPLATE_PATH.read_bytes()).hexdigest()[:12]
TEMPLATE_NOTICE = (
    "原模板历史主体、部门样例和外部引用已清除，失效明细合计按实际员工范围修复。"
    "保留原工作表与项目区域，项目工时、税后调整、税务和其他无来源字段留空，不作项目分摊。"
    "主体按 ID 顺序使用模板页，页内公司名称为实际主体；未使用页保持空模板。"
    "部门层级无历史快照，汇总仅按主体，明细保留锁定部门名称；公式需 Excel 打开后重算。"
)
SOCIAL_COLUMNS = {
    "AP": "social_employee_pension",
    "AQ": "social_employee_medical",
    "AR": "social_employee_unemployment",
    "AY": "social_company_pension",
    "AZ": "social_company_medical",
    "BA": "social_company_unemployment",
    "BB": "social_company_injury",
    "BC": "social_company_maternity",
}


def render_labor_cost(ledger: dict) -> bytes:
    rows = sorted(
        ledger["records"], key=lambda r: (r["subject_id"], r["employee_id"], r["payroll_batch_id"])
    )
    if not rows:
        raise ValueError("无工资记录，不能生成人工成本表")
    book = load_workbook(TEMPLATE_PATH)
    detail = book["工资明细"]
    detail["AD1"] = f"{ledger['period']} 人工成本表"
    detail["H3"] = "部门（锁定快照）"
    for n, record in enumerate(rows, 5):
        for col in range(1, 69):
            detail.cell(n, col)._style = copy(detail.cell(4, col)._style)
        snapshot, amounts = record["snapshot"], record["amounts"]
        items = {item["code"]: Decimal(item["amount"]) for item in record["items"]}
        source = "独立补发" if record["batch_type"] == "supplement" else "正常工资"
        source += f" #{record['payroll_batch_id']}（批次号 {record['batch_no']}）"
        if record["batch_type"] == "supplement":
            source += f"：{record['batch_name']}"
        if record.get("correction_of_batch_id"):
            source += f"；整批更正替代 #{record['correction_of_batch_id']}"
        values = {
            "A": n - 4,
            "B": record["subject_name"],
            "C": snapshot["name"],
            "D": snapshot["id_number"],
            "H": snapshot.get("department_name"),
            "K": snapshot.get("position_title"),
            "M": snapshot.get("level_code"),
            "AO": Decimal(amounts["gross"]),
            "AS": sum(
                (v for k, v in items.items() if k.startswith("social_employee_")), Decimal(0)
            ),
            "AT": items.get("housing_employee", Decimal(0)),
            "AX": Decimal(amounts["untaxed_amount"]),
            "BD": sum((v for k, v in items.items() if k.startswith("social_company_")), Decimal(0)),
            "BE": items.get("housing_company", Decimal(0)),
            "BF": Decimal(amounts["employer_cost"]),
            "BG": source + "；未扣个税金额",
            "BM": record["subject_name"],
            "BN": snapshot.get("bank_account"),
            **{col: items.get(code) for col, code in SOCIAL_COLUMNS.items()},
        }
        for col, value in values.items():
            _write(detail, n, col, value)
    last = len(rows) + 4
    detail["B4"] = "合计"
    for col in (*SOCIAL_COLUMNS, "AO", "AS", "AT", "AX", "BD", "BE", "BF"):
        source_range = f"{col}5:{col}{last}"
        detail[f"{col}4"] = (
            f'=IF(COUNT({source_range})=0,"",SUM({source_range}))'
            if col in SOCIAL_COLUMNS
            else f"=SUM({source_range})"
        )
        detail[f"{col}4"].number_format = "#,##0.00"
    detail.freeze_panes = "E5"
    detail.auto_filter.ref = f"A3:BP{last}"
    detail.print_area = f"A1:BP{last}"

    templates = book.worksheets[:-1]
    subjects = sorted({r["subject_id"] for r in rows})
    while len(templates) < len(subjects):
        sheet = book.copy_worksheet(templates[0])
        sheet.title = f"主体{subjects[len(templates)]}人力成本"
        templates.append(sheet)
    for index, subject_id in enumerate(subjects):
        sheet = templates[index]
        sheet.sheet_state = "visible"
        indices = [i for i, r in enumerate(rows, 5) if r["subject_id"] == subject_id]
        records = [r for r in rows if r["subject_id"] == subject_id]
        _write(sheet, 1, "B", ledger["period"])
        _write(sheet, 1, "F", records[0]["subject_name"])
        sheet["A2"] = "公司成本（未扣个税口径）"
        sheet["E2"] = f"=SUM('工资明细'!BF{indices[0]}:BF{indices[-1]})"
        sheet["E2"].number_format = "#,##0.00"
        headers = [r for r in range(1, sheet.max_row + 1) if sheet.cell(r, 1).value == "一级部门"]
        maps = [
            {"E": "AO", "F": "AP", "G": "AQ", "H": "AR", "I": "AS", "J": "AT", "N": "AX"},
            {"E": "BE", "F": "AT"},
            {"E": "AY", "F": "AZ", "G": "BA", "H": "BB", "I": "BC", "J": "BD", "K": "AS"},
        ]
        for header, mapping in zip(headers, maps, strict=True):
            n = header + 2
            sheet[f"A{n}"] = "主体汇总"
            sheet[f"D{n}"] = len({r["employee_id"] for r in records})
            for col, target in mapping.items():
                source_range = f"'工资明细'!{target}{indices[0]}:{target}{indices[-1]}"
                sheet[f"{col}{n}"] = (
                    f'=IF(COUNT({source_range})=0,"",SUM({source_range}))'
                    if target in SOCIAL_COLUMNS
                    else f"=SUM({source_range})"
                )
                sheet[f"{col}{n}"].number_format = "#,##0.00"
                if target in SOCIAL_COLUMNS:
                    total = next(
                        r
                        for r in range(n + 1, sheet.max_row + 1)
                        if sheet.cell(r, 1).value == "合计"
                    )
                    body_range = f"{col}{n}:{col}{total - 1}"
                    sheet[f"{col}{total}"] = f'=IF(COUNT({body_range})=0,"",SUM({body_range}))'
            if header == headers[1]:
                sheet[f"G{n}"] = f"=SUM(E{n}:F{n})"
            elif header == headers[2]:
                sheet[f"L{n}"] = f"=SUM(J{n}:K{n})"
        last_total = max(r for r in range(1, sheet.max_row + 1) if sheet.cell(r, 1).value == "合计")
        sheet.print_area = f"A1:N{last_total}"
    output = BytesIO()
    book.save(output)
    return output.getvalue()
