"""工资表模板适配器：只使用已锁定台账快照，不重新算薪或计算个税。"""

from copy import copy
from decimal import Decimal
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Alignment

TEMPLATE_PATH = Path(__file__).with_name("templates") / "payroll-v1.xlsx"
TEMPLATE_VERSION = "payroll-v1-" + sha256(TEMPLATE_PATH.read_bytes()).hexdigest()[:12]
TAX_NOTICE = "未扣个税金额；个税及税务申报字段无来源，已留空，请在 Excel 线下计算并核对后发薪。"
TEMPLATE_NOTICE = (
    "原模板历史样例已清除；工资合计原公式仅引用表头，已修正为实际员工范围。"
    "税务计算、代发、研发工时等辅助页保留结构但不填充；无来源字段留空。"
    "汇总公式请用 Excel 打开后重算，员工金额直接取锁定台账。"
)


def _write(sheet, row: int, column: str, value: Any) -> None:
    cell = sheet[f"{column}{row}"]
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"  # Names and identifiers must never become Excel formulas.
        cell.number_format = "@"
    elif isinstance(value, Decimal):
        cell.number_format = "#,##0.00"


def render_payroll(ledger: dict[str, Any]) -> bytes:
    book = load_workbook(TEMPLATE_PATH)
    detail = book["工资表明细"]
    rows = sorted(
        ledger["records"], key=lambda row: (row["subject_id"], row["employee_id"], row["batch_no"])
    )
    last = len(rows) + 4
    if not rows:
        raise ValueError("无工资记录，不能生成工资表")
    detail["AD1"] = f"{ledger['period']} 工资表"
    detail["B4"] = "合计"
    # The template has no employee row. Reuse its styled numeric total row as body style.
    for n, record in enumerate(rows, 5):
        for col in range(1, 69):
            cell = detail.cell(n, col)
            cell._style = copy(detail.cell(4, col)._style)
            cell.alignment = Alignment(vertical="center")
        snapshot = record["snapshot"]
        amounts = record["amounts"]
        items = {i["code"]: Decimal(i["amount"]) for i in record["items"]}

        def total(prefix):
            return sum((v for k, v in items.items() if k.startswith(prefix)), Decimal(0))

        incentive = items.get("attendance_incentive", Decimal(0))
        attendance = total("attendance_") - incentive
        source = (
            f"独立补发 S{record['batch_no']}：{record['batch_name']}"
            if record["batch_type"] == "supplement"
            else f"正常工资 N{record['batch_no']}"
        )
        if record.get("correction_of_batch_id"):
            source += f"；整批更正，替代批次 #{record['correction_of_batch_id']}"
        values = {
            "A": n - 4,
            "B": record["subject_name"],
            "C": snapshot["name"],
            "D": snapshot["id_number"],
            "H": snapshot.get("department_name"),
            "K": snapshot.get("position_title"),
            "M": snapshot.get("level_code"),
            "AO": Decimal(amounts["gross"]),
            "AX": Decimal(amounts["untaxed_amount"]),
            "BF": Decimal(amounts["employer_cost"]),
            "BG": source + "；未扣个税",
            "BM": record["subject_name"],
            "BN": snapshot.get("bank_account"),
        }
        if record["batch_type"] == "supplement":
            values["AM"] = Decimal(amounts["gross"])
        else:
            values.update(
                {
                    "O": Decimal(snapshot["fixed_salary"]) + Decimal(snapshot["performance_base"]),
                    "P": Decimal(snapshot["fixed_salary"]),
                    "AC": Decimal(snapshot["performance_base"]),
                    "Q": items.get("attendance_late", Decimal(0))
                    + items.get("attendance_early", Decimal(0)),
                    "R": items.get("attendance_missed_punch", Decimal(0)),
                    "X": attendance,
                    "Y": items.get("attendance_incentive", Decimal(0)),
                    "AB": items.get("fixed_salary", Decimal(0)) - attendance + incentive,
                    "AJ": items.get("performance", Decimal(0)),
                    "AS": total("social_employee_"),
                    "AT": items.get("housing_employee", Decimal(0)),
                    "BD": total("social_company_"),
                    "BE": items.get("housing_company", Decimal(0)),
                }
            )
            values["BG"] += "；无薪请假扣款计入考勤合计，不猜测假别"
        for col, value in values.items():
            _write(detail, n, col, value)
        slip = book["工资条"]
        for col, value in {
            "A": snapshot["id_number"],
            "B": snapshot["name"],
            "O": values["AO"],
            "P": values.get("AS"),
            "Q": values.get("AT"),
            "U": values["AX"],
            "V": values["BG"],
        }.items():
            _write(slip, n - 3, col, value)
    for col in (
        "O",
        "P",
        "Q",
        "R",
        "X",
        "Y",
        "AB",
        "AC",
        "AJ",
        "AM",
        "AO",
        "AS",
        "AT",
        "AX",
        "BD",
        "BE",
        "BF",
    ):
        detail[f"{col}4"] = f"=SUBTOTAL(9,{col}5:{col}{last})"
        detail[f"{col}4"].number_format = "#,##0.00"
    detail.freeze_panes = "E5"
    detail.auto_filter.ref = f"A3:BP{last}"
    detail.print_area = f"A1:BP{last}"
    summary = book["汇总"]
    summary["B1"] = f"{ledger['period']} 工资情况（未扣个税金额）"
    subjects = {r["subject_id"]: r["subject_name"] for r in rows}
    for n, (subject_id, name) in enumerate(subjects.items(), 3):
        for col in range(2, 10):
            summary.cell(n, col)._style = copy(summary.cell(3, col)._style)
        _write(summary, n, "B", name)
        indices = [i for i, row in enumerate(rows, 5) if row["subject_id"] == subject_id]
        for col, target in {"C": "AX", "E": "BD", "F": "AS", "G": "BE", "H": "AT"}.items():
            summary[f"{col}{n}"] = f"=SUM('工资表明细'!{target}{indices[0]}:{target}{indices[-1]})"
            summary[f"{col}{n}"].number_format = "#,##0.00"
        summary[f"I{n}"] = len({r["employee_id"] for r in rows if r["subject_id"] == subject_id})
    end = len(subjects) + 3
    summary[f"B{end}"] = "总计"
    for col in ("C", "E", "F", "G", "H", "I"):
        summary[f"{col}{end}"] = f"=SUM({col}3:{col}{end - 1})"
    summary.print_area = f"B1:I{end}"
    output = BytesIO()
    book.save(output)
    return output.getvalue()
