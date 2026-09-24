"""Fixed monthly performance coefficient workbook."""

from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from paylite.excel.employee_template import MAX_ROWS, WorkbookFormatError, WorkbookRow, _value

TEMPLATE_VERSION = "TPL-PERF-v1"
SHEET_NAME = "月度绩效"
HEADERS = ("身份证号", "姓名", "绩效系数", "来源说明")


def make_template() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    sheet.append([TEMPLATE_VERSION])
    sheet.append(HEADERS)
    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = "A2:D2"
    for cell in sheet[2]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="315FD1")
        cell.alignment = Alignment(vertical="center")
        sheet.column_dimensions[cell.column_letter].width = 27 if cell.column == 1 else 22
    sheet.column_dimensions["D"].width = 36
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def parse_workbook(content: bytes) -> list[WorkbookRow]:
    try:
        workbook = load_workbook(
            BytesIO(content), read_only=False, data_only=False, keep_links=False
        )
    except Exception as exc:
        raise WorkbookFormatError("Excel 文件无法读取") from exc
    try:
        if workbook.sheetnames != [SHEET_NAME]:
            raise WorkbookFormatError("工作表不匹配，请使用月度绩效固定模板")
        sheet = workbook[SHEET_NAME]
        if sheet["A1"].value != TEMPLATE_VERSION:
            raise WorkbookFormatError(f"模板版本不匹配，需要 {TEMPLATE_VERSION}")
        if (
            sheet.max_column != len(HEADERS)
            or tuple(sheet.cell(2, col).value for col in range(1, len(HEADERS) + 1)) != HEADERS
        ):
            raise WorkbookFormatError("模板表头或列顺序不匹配")
        if sheet.max_row - 2 > MAX_ROWS:
            raise WorkbookFormatError(f"单次最多导入 {MAX_ROWS} 行")
        rows = []
        for number in range(3, sheet.max_row + 1):
            raw = {}
            errors = []
            for column, field in enumerate(HEADERS, 1):
                value, error = _value(sheet.cell(number, column), field)
                raw[field] = value
                if error:
                    errors.append(error)
            if any(raw.values()):
                rows.append(WorkbookRow(number, raw, errors))
        if not rows:
            raise WorkbookFormatError("Excel 文件没有绩效数据行")
        return rows
    finally:
        workbook.close()
