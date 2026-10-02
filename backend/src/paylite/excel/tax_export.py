"""申报辅助模板只填已锁定事实，不计算个税或推断申报口径。"""

from copy import copy
from decimal import Decimal
from hashlib import sha256
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from paylite.excel.payroll_export import _write

TEMPLATE_PATH = Path(__file__).with_name("templates") / "tax-v1.xlsx"
TEMPLATE_VERSION = "tax-v1-" + sha256(TEMPLATE_PATH.read_bytes()).hexdigest()[:12]
TEMPLATE_NOTICE = (
    "旧式 xls 已通过 LibreOffice 转换为 xlsx，保留 Sheet1 的 33 列表头和样式。"
    "本次原件核对未发现公式、外部引用或错误示例；没有添加税额计算公式。"
    "本期收入取锁定应发金额，各有效批次单列来源；未扣个税金额另在备注标识。"
    "所得期间及证件类型无申报来源留空，不按工资期间或发放日期猜测。"
    "仅供线下补充核对，未完成税务申报；不回填系统。"
)
SOCIAL_COLUMNS = {
    "I": "social_employee_pension",
    "J": "social_employee_medical",
    "K": "social_employee_unemployment",
    "L": "housing_employee",
}
BLANK_FIELDS = {
    "C": "证件类型",
    "E": "所得期间起",
    "F": "所得期间止",
    "H": "本期免税收入",
    "M": "累计子女教育",
    "N": "累计继续教育",
    "O": "累计住房贷款利息",
    "P": "累计住房租金",
    "Q": "累计赡养老人",
    "R": "累计3岁以下婴幼儿照护",
    "S": "累计个人养老金",
    "T": "企业(职业)年金",
    "U": "商业健康保险",
    "V": "税延养老保险",
    "W": "公务交通费用",
    "X": "通讯费用",
    "Y": "律师办案费用",
    "Z": "住房公积金调整",
    "AA": "准予扣除的捐赠额",
    "AB": "税前扣除项目合计",
    "AC": "减免税额",
    "AD": "协定减免",
    "AE": "减除费用标准",
    "AF": "已缴税额",
}


def render_tax(ledger: dict) -> bytes:
    book = load_workbook(TEMPLATE_PATH)
    sheet = book["Sheet1"]
    rows = sorted(
        ledger["records"],
        key=lambda r: (r["subject_id"], r["employee_id"], r["payroll_batch_id"]),
    )
    if not rows:
        raise ValueError("无工资记录，不能生成申报辅助模板")
    for n, record in enumerate(rows, 2):
        for col in range(1, 34):
            sheet.cell(n, col)._style = copy(sheet.cell(1, col)._style)
        snapshot = record["snapshot"]
        source = "独立补发" if record["batch_type"] == "supplement" else "正常工资"
        source += f" #{record['payroll_batch_id']}（批次号 {record['batch_no']}）"
        if record["batch_type"] == "supplement":
            source += f"：{record['batch_name']}"
        if record.get("correction_of_batch_id"):
            source += f"；整批更正替代 #{record['correction_of_batch_id']}"
        remark = (
            f"{record['subject_name']}；工资期间 {ledger['period']}；{source}；"
            f"未扣个税金额 {record['amounts']['untaxed_amount']}；"
            "仅供线下补充核对，未完成税务申报"
        )
        if len(remark) > 32767:
            raise ValueError("来源备注超过 Excel 单元格上限，不能截断导出")
        values = {
            "A": snapshot.get("employee_no"),
            "B": snapshot["name"],
            "D": snapshot["id_number"],
            "G": Decimal(record["amounts"]["gross"]),
            "AG": remark,
        }
        items = {item["code"]: Decimal(item["amount"]) for item in record["items"]}
        values.update({col: items.get(code) for col, code in SOCIAL_COLUMNS.items()})
        for col, value in values.items():
            _write(sheet, n, col, value)
    sheet.freeze_panes = "E2"
    sheet.auto_filter.ref = f"A1:AG{len(rows) + 1}"
    sheet.oddFooter.center.text = "未扣个税金额；申报辅助，未完成税务申报"
    output = BytesIO()
    book.save(output)
    return output.getvalue()
