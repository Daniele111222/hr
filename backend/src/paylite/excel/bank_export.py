"""将已校验、合并的代发记录写入显式选择的主体模板。"""

from copy import copy
from decimal import Decimal
from hashlib import sha256
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from paylite.excel.payroll_export import _write

TEMPLATE_PATH = Path(__file__).with_name("templates") / "bank-v1.xlsx"
TEMPLATE_VERSION = "bank-v1-" + sha256(TEMPLATE_PATH.read_bytes()).hexdigest()[:12]
TEMPLATE_SHEETS = tuple(load_workbook(TEMPLATE_PATH).sheetnames)
TEMPLATE_NOTICE = (
    "原模板历史样例与 #REF! 公式已清除，合计公式按实际员工范围生成。"
    "开户行无锁定快照来源，留空待线下核对；户名取锁定员工姓名。"
    "未选用工作表保留空模板；汇总公式需 Excel 打开后重算。"
)


def render_bank(rows: list[dict], subject_templates: dict[str, str]) -> bytes:
    if not rows:
        raise ValueError("无代发记录")
    book = load_workbook(TEMPLATE_PATH)
    used = set()
    for subject_id in sorted({r["subject_id"] for r in rows}):
        title = subject_templates[str(subject_id)]
        sheet = book[title]
        if title in used:
            sheet = book.copy_worksheet(sheet)
            sheet.title = f"{title}-S{subject_id}"
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    cell.value = None
        used.add(title)
        columns = {cell.value: cell.column_letter for cell in sheet[1] if cell.value}
        records = [r for r in rows if r["subject_id"] == subject_id]
        for n, record in enumerate(records, 2):
            for col in range(1, sheet.max_column + 1):
                sheet.cell(n, col)._style = copy(sheet.cell(2, col)._style)
            for label, value in {
                "编号": n - 1,
                "员工账号": record["bank_account"],
                "姓名": record["name"],
                "身份证号": record["id_number"],
                "未扣个税金额": Decimal(record["amount"]),
                "备注": record["remark"],
            }.items():
                if label in columns:
                    _write(sheet, n, columns[label], value)
        end = len(records) + 2
        amount = columns["未扣个税金额"]
        sheet[f"{columns['姓名']}{end}"] = "合计"
        sheet[f"{amount}{end}"] = f"=SUM({amount}2:{amount}{end - 1})"
        sheet[f"{amount}{end}"].number_format = "#,##0.00"
        sheet.print_area = f"A1:{columns['备注']}{end}"
    output = BytesIO()
    book.save(output)
    return output.getvalue()
