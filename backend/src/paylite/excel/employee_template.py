"""Fixed employee workbook format. No business decisions live in this adapter."""

from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Font, PatternFill

TEMPLATE_VERSION = "TPL-EMP-v2.4"
SHEET_NAME = "员工资料"
MAX_FILE_SIZE = 20 * 1024 * 1024
MAX_ROWS = 5000
HEADERS = (
    "身份证号",
    "员工编号",
    "姓名",
    "员工类型",
    "入职日期",
    "离职日期",
    "转正状态",
    "转正日期",
    "是否正式",
    "主体编码",
    "部门编码",
    "职位",
    "职级",
    "职级数字",
    "固定薪资",
    "绩效基数",
    "base城市编码",
    "银行卡号",
    "开户名",
    "银行名称",
    "支行名称",
    "生效日期",
)
TEXT_IDENTIFIERS = {"身份证号", "员工编号", "银行卡号"}


class WorkbookFormatError(ValueError):
    pass


@dataclass(frozen=True)
class WorkbookRow:
    number: int
    raw: dict[str, str]
    errors: list[dict[str, str]]


def make_template() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    sheet.append([TEMPLATE_VERSION])
    sheet.append(HEADERS)
    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = "A2:V2"
    for cell in sheet[2]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="315FD1")
        cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[2].height = 24
    for index, header in enumerate(HEADERS, 1):
        letter = sheet.cell(2, index).column_letter
        sheet.column_dimensions[letter].width = 21 if header not in TEXT_IDENTIFIERS else 27
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _value(cell: Cell, field: str) -> tuple[str, dict[str, str] | None]:
    if cell.value is None:
        return "", None
    if cell.data_type == "f":
        return str(cell.value), {
            "code": "FORMULA_NOT_ALLOWED",
            "field": field,
            "column": cell.column_letter,
            "cell": cell.coordinate,
            "message": "导入模板不接受公式，请填写实际值",
        }
    if cell.data_type == "e":
        return str(cell.value), {
            "code": "EXCEL_ERROR",
            "field": field,
            "column": cell.column_letter,
            "cell": cell.coordinate,
            "message": "Excel 单元格含错误值",
        }
    if field in TEXT_IDENTIFIERS and cell.data_type != "s":
        return str(cell.value), {
            "code": "IDENTIFIER_NOT_TEXT",
            "field": field,
            "column": cell.column_letter,
            "cell": cell.coordinate,
            "message": f"{field}必须按文本填写，避免前导零或长数字丢失",
        }
    value: Any = cell.value
    if isinstance(value, (date, datetime)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat(), None
    return str(value).strip(), None


def parse_workbook(content: bytes) -> list[WorkbookRow]:
    try:
        workbook = load_workbook(
            BytesIO(content), read_only=False, data_only=False, keep_links=False
        )
    except Exception as exc:
        raise WorkbookFormatError("Excel 文件无法读取") from exc
    try:
        if workbook.sheetnames != [SHEET_NAME]:
            raise WorkbookFormatError("工作表不匹配，请使用员工资料固定模板")
        sheet = workbook[SHEET_NAME]
        if sheet["A1"].value != TEMPLATE_VERSION:
            raise WorkbookFormatError(f"模板版本不匹配，需要 {TEMPLATE_VERSION}")
        if tuple(sheet.cell(2, col).value for col in range(1, len(HEADERS) + 1)) != HEADERS:
            raise WorkbookFormatError("模板表头或列顺序不匹配")
        if sheet.max_row - 2 > MAX_ROWS:
            raise WorkbookFormatError(f"单次最多导入 {MAX_ROWS} 行")
        rows = []
        for number in range(3, sheet.max_row + 1):
            raw: dict[str, str] = {}
            errors = []
            for column, field in enumerate(HEADERS, 1):
                value, error = _value(sheet.cell(number, column), field)
                raw[field] = value
                if error:
                    errors.append(error)
            if any(raw.values()):
                rows.append(WorkbookRow(number, raw, errors))
        if not rows:
            raise WorkbookFormatError("Excel 文件没有员工数据行")
        return rows
    finally:
        workbook.close()
